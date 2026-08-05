#!/usr/bin/env python3
"""
doc_extract.py — Extracteur Markdown dédié au manuel FRED DNB14A.

Utilisation :
    python tools/doc_extract.py A
    python tools/doc_extract.py JM
    python tools/doc_extract.py ZD

Entrée :
    archive/DNB14A.docx

Sortie :
    docs/fr/reference/commandes/<commande>.md

Dépendance :
    python-docx
"""

from __future__ import annotations

import re
import sys
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
except ImportError:
    print(
        "Erreur : le module 'python-docx' n'est pas installé.\n"
        "Installe-le avec : python -m pip install python-docx",
        file=sys.stderr,
    )
    raise SystemExit(2)


# ---------------------------------------------------------------------------
# Configuration fixe du projet
# ---------------------------------------------------------------------------

SCRIPT_PATH = Path(__file__).resolve()
REPO_ROOT = SCRIPT_PATH.parent.parent

SOURCE_CANDIDATES = (
    REPO_ROOT / "archive" / "DNB14A.docx",
    REPO_ROOT / "DNB14A.docx",
)

OUTPUT_DIR = REPO_ROOT / "docs" / "fr" / "reference" / "commandes"

SECTION_LABELS = {
    "syntaxe": "syntax",
    "exemples": "examples",
    "descriptions": "description",
    "description": "description",
    "options": "options",
    "remarques": "remarks",
    "remarque": "remarks",
}

# Les alias évitent de saisir certains caractères délicats dans le terminal.
COMMAND_ALIASES = {
    "FQ": "F?",
    "BANG": "!",
    "QUOTE": '"',
    "COUNTER": "#",
    "EQUALS": "=",
    "COLON": ":",
    "AT": "@",
    "BACKSLASH": "\\\\",
}

# Noms de fichiers sûrs pour les commandes spéciales.
FILE_NAMES = {
    "F?": "fq.md",
    "!": "bang.md",
    '"': "quote.md",
    "#": "counter.md",
    "=": "equals.md",
    ":": "colon.md",
    "@": "at.md",
    "\\\\": "backslashbackslash.md",
    "N:": "ncolon.md",
    "N=": "nequals.md",
    "N<": "nless.md",
    "N>": "ngreater.md",
    "N+": "nplus.md",
    "N-": "nminus.md",
    "N*": "nmultiply.md",
    "N/": "ndivide.md",
    "N%": "nmodulo.md",
    "N~": "ncomplement.md",
    "N&": "nand.md",
    "N^": "nxor.md",
    "N{": "nshiftleft.md",
    "N}": "nshiftright.md",
    "N|": "npipe.md",
    "OI(": "oi_parenthesis.md",
    "OI\\F": "oi_backslash_f.md",
    "OS&": "os_ampersand.md",
    "OS/": "os_separator.md",
}

# Titres harmonisés lorsqu'une commande possède plusieurs entrées dans DNB14A.
TITLE_OVERRIDES = {
    "F": "F — Informations sur le buffer et association de fichier (Facts / File)",
    "J": "J — Sauts et branchements (Jump)",
    "F?": "F? — Liste des buffers modifiés",
    "\\\\": "\\\\ — Attente",
}

# Descriptions courtes utilisées dans le tableau Paramètres.
PARAMETER_DESCRIPTIONS = {
    "<bl>": "Un espace.",
    "<texte>": "Texte ou lignes de texte traités par la commande.",
    "<nom de buffer>": "Nom du buffer concerné.",
    "<nom buffer>": "Nom du buffer concerné.",
    "<nom de modèle>": "Nom attribué au modèle.",
    "<modèle>": "Modèle de recherche FRED.",
    "<commandes>": "Commandes FRED à exécuter.",
    "<commande>": "Commande FRED à exécuter.",
    "<nom de fichier>": "Nom du fichier concerné.",
    "<liste fichiers>": "Liste de fichiers examinée dans l'ordre.",
    "<nom de registre>": "Nom du registre numérique.",
    "<nom registre>": "Nom du registre numérique.",
    "<nombre>": "Valeur numérique.",
    "<caractère>": "Caractère utilisé par la commande ou l'option.",
    "<caractères>": "Ensemble de caractères.",
    "<chaîne>": "Chaîne de caractères.",
    "<chaîne1>": "Première chaîne de caractères.",
    "<chaîne2>": "Deuxième chaîne de caractères.",
    "<remarques>": "Message à afficher.",
    "<label>": "Nom d'étiquette.",
    "<n>": "Nombre entier.",
    "<nl>": "Fin de ligne.",
    "<cr>": "Retour chariot.",
}

SEE_ALSO = {
    "A": ("I", "C", "D"),
    "B": ("FB", "F", "ZD"),
    "C": ("A", "I", "D"),
    "D": ("A", "C", "V"),
    "E": ("FE", "S", "T"),
    "F": ("FB", "FF", "F?"),
    "FB": ("B", "F", "FF"),
    "FE": ("E", "T", "S"),
    "FF": ("F", "FB", "F?"),
    "FN": ("N:", "NP"),
    "FO": ("O", "FB", "FN"),
    "FV": ("FO",),
    "F?": ("F", "FB", "FF"),
    "G": ("T", "U", "J"),
    "I": ("A", "C", "D"),
    "J": ("JM", "JP", "JE", "JO", "JB"),
    "JM": ("JP", "JE", "J"),
    "JP": ("JM", "JE", "J"),
    "JE": ("JM", "JP", "OQ"),
    "JO": ("J", "JB", "G", "U"),
    "JB": ("J", "JO", "U"),
    "K": ("ZK", "M", "ZM"),
    "L": ("P", "R", "W"),
    "M": ("K", "ZK", "ZM"),
}


# ---------------------------------------------------------------------------
# Lecture et détection des sections
# ---------------------------------------------------------------------------

def locate_source() -> Path:
    """Retourne le chemin du DNB14A.docx."""
    for candidate in SOURCE_CANDIDATES:
        if candidate.is_file():
            return candidate

    expected = "\n".join(f"  - {path}" for path in SOURCE_CANDIDATES)
    raise FileNotFoundError(
        "DNB14A.docx est introuvable. Emplacements recherchés :\n" + expected
    )


def normalize_command(value: str) -> str:
    """Normalise la commande fournie sur la ligne de commande."""
    value = value.strip()
    if not value:
        raise ValueError("La commande ne peut pas être vide.")

    upper = value.upper()
    return COMMAND_ALIASES.get(upper, upper)


def load_paragraphs(path: Path) -> list[str]:
    """Charge les paragraphes non vides du document."""
    document = Document(path)
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def split_heading(text: str) -> tuple[str, str] | None:
    """
    Reconnaît un vrai titre de commande.

    Les lignes du sommaire contiennent une tabulation et un numéro de page :
    elles sont volontairement ignorées.
    """
    if "\t" in text:
        return None

    match = re.match(r"^(.+?)\s+-\s+(.+?)\s*$", text)
    if not match:
        return None

    command = match.group(1).strip()
    title = match.group(2).strip()

    # Les marqueurs de section ne sont pas des commandes.
    if command.casefold() in SECTION_LABELS:
        return None

    return command, title


def collect_command_sections(paragraphs: list[str]) -> dict[str, list[tuple[str, list[str]]]]:
    """
    Construit un index :
        commande -> [(titre, paragraphes), ...]

    Plusieurs entrées portant le même nom sont conservées. C'est nécessaire
    pour F et J, qui possèdent plusieurs formes documentées.
    """
    sections: dict[str, list[tuple[str, list[str]]]] = {}
    current_command: str | None = None
    current_title = ""
    current_body: list[str] = []
    in_reference = False

    def flush() -> None:
        nonlocal current_command, current_title, current_body
        if current_command is not None:
            sections.setdefault(current_command, []).append(
                (current_title, current_body.copy())
            )
        current_command = None
        current_title = ""
        current_body = []

    for paragraph in paragraphs:
        heading = split_heading(paragraph)

        # Le premier vrai titre de commande après le sommaire est A.
        if not in_reference:
            if heading and heading[0] == "A":
                in_reference = True
            else:
                continue

        if heading:
            flush()
            current_command, current_title = heading
        elif current_command is not None:
            # Le pied de page historique n'appartient pas à la dernière commande.
            if paragraph.startswith("Bull S.A.") or paragraph.startswith("University of Waterloo"):
                continue
            current_body.append(paragraph)

    flush()
    return sections


# ---------------------------------------------------------------------------
# Analyse du contenu d'une commande
# ---------------------------------------------------------------------------

def classify_body(body: Iterable[str]) -> dict[str, list[str]]:
    """Répartit les paragraphes dans Syntaxe, Exemples, Description, etc."""
    result = {
        "syntax": [],
        "examples": [],
        "description": [],
        "options": [],
        "remarks": [],
        "other": [],
    }
    current = "other"

    for paragraph in body:
        key = paragraph.rstrip(":").strip().casefold()
        if key in SECTION_LABELS:
            current = SECTION_LABELS[key]
            continue
        result[current].append(paragraph)

    return result


def clean_text(paragraphs: Iterable[str]) -> str:
    """Nettoie sans réécrire le contenu historique."""
    text = "\n\n".join(item.strip() for item in paragraphs if item.strip())
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_code(paragraphs: Iterable[str]) -> str:
    """
    Préserve les exemples historiques dans un bloc de code.

    Le DOCX converti depuis Word 2.0 a parfois concaténé plusieurs sorties.
    La V1 reste volontairement conservatrice : elle ne devine pas les
    séparations qui ne sont plus présentes dans la source.
    """
    lines = [item.rstrip() for item in paragraphs if item.strip()]
    return "\n".join(lines).strip()


def first_sentence(text: str) -> str:
    """Extrait une phrase de résumé depuis la description."""
    compact = re.sub(r"\s+", " ", text).strip()
    if not compact:
        return "Documentation de la commande FRED."

    match = re.match(r"(.+?[.!?])(?:\s|$)", compact)
    sentence = match.group(1) if match else compact

    if len(sentence) > 220:
        sentence = sentence[:217].rstrip() + "..."
    return sentence


def extract_parameters(syntax: str) -> list[tuple[str, str]]:
    """Déduit un petit tableau de paramètres depuis la syntaxe."""
    parameters: list[tuple[str, str]] = []

    # Préfixes d'adresses les plus fréquents.
    address_prefixes = re.findall(r"\([^)\n]+\)", syntax)
    for prefix in address_prefixes:
        if prefix not in {value for value, _ in parameters}:
            parameters.append(
                (
                    prefix,
                    "Adresse ou plage de lignes optionnelle, selon la forme indiquée.",
                )
            )

    placeholders = re.findall(r"<[^>\n]+>", syntax)
    for placeholder in placeholders:
        if placeholder in {value for value, _ in parameters}:
            continue
        description = PARAMETER_DESCRIPTIONS.get(
            placeholder,
            "Élément défini par la syntaxe historique de la commande.",
        )
        parameters.append((placeholder, description))

    return parameters


def filename_for(command: str) -> str:
    """Retourne le nom de fichier Markdown associé à la commande."""
    if command in FILE_NAMES:
        return FILE_NAMES[command]

    safe = command.casefold()
    safe = safe.replace("\\", "_backslash_")
    safe = re.sub(r"[^a-z0-9_]+", "_", safe)
    safe = safe.strip("_")
    if not safe:
        raise ValueError(f"Impossible de construire un nom de fichier pour {command!r}.")
    return safe + ".md"


def link_for(command: str) -> str:
    """Construit un lien Markdown vers une autre commande."""
    return filename_for(command)


def build_title(command: str, titles: list[str]) -> str:
    """Construit le titre principal de la page."""
    if command in TITLE_OVERRIDES:
        return TITLE_OVERRIDES[command]

    title = titles[0].rstrip(".") if titles else "Commande FRED"
    return f"{command} — {title}"


def build_markdown(
    command: str,
    source_sections: list[tuple[str, list[str]]],
) -> str:
    """Produit le Markdown selon le gabarit validé du Lot 1."""
    titles: list[str] = []
    syntaxes: list[str] = []
    examples: list[str] = []
    descriptions: list[str] = []
    options: list[str] = []
    remarks: list[str] = []

    for title, body in source_sections:
        titles.append(title)
        classified = classify_body(body)

        syntax = clean_code(classified["syntax"])
        example = clean_code(classified["examples"])
        description = clean_text(classified["description"])
        option = clean_text(classified["options"])
        remark = clean_text(classified["remarks"])
        other = clean_text(classified["other"])

        if syntax:
            syntaxes.append(syntax)
        if example:
            examples.append(example)
        if description:
            descriptions.append(description)
        elif other:
            descriptions.append(other)
        if option:
            options.append(option)
        if remark:
            remarks.append(remark)

    title = build_title(command, titles)
    description_text = "\n\n".join(descriptions).strip()
    summary = first_sentence(description_text)
    syntax_text = "\n\n".join(syntaxes).strip()
    examples_text = "\n\n".join(examples).strip()
    options_text = "\n\n".join(options).strip()
    remarks_text = "\n\n".join(remarks).strip()
    parameters = extract_parameters(syntax_text)

    output: list[str] = [
        f"# {title}",
        "",
        f"> {summary}",
        "",
        "## Syntaxe",
        "",
        "```fred",
        syntax_text or command,
        "```",
        "",
        "## Paramètres",
        "",
    ]

    if parameters:
        output.extend(
            [
                "| Élément | Description |",
                "|---|---|",
            ]
        )
        for element, meaning in parameters:
            output.append(f"| `{element}` | {meaning} |")
    else:
        output.append("Aucun paramètre explicite n'est documenté.")

    output.extend(
        [
            "",
            "## Description",
            "",
            description_text or "Aucune description n'est disponible dans la source.",
            "",
            "## Exemples",
            "",
        ]
    )

    if examples_text:
        output.extend(["```fred", examples_text, "```"])
    else:
        output.append("Aucun exemple n'est fourni dans la source.")

    output.extend(
        [
            "",
            "## Options",
            "",
            options_text or "Aucune option spécifique n'est documentée.",
            "",
            "## Remarques",
            "",
            remarks_text or "Aucune remarque spécifique n'est documentée.",
            "",
            "## Compatibilité",
            "",
            "| Implémentation | Statut |",
            "|---|---|",
            "| FRED historique | Compatible |",
            "| FRED++ | À valider selon l'implémentation |",
            "",
            "## Voir aussi",
            "",
        ]
    )

    related = SEE_ALSO.get(command, ())
    if related:
        for item in related:
            output.append(f"- [{item}]({link_for(item)})")
    else:
        output.append("Aucun lien associé n'est encore défini.")

    return "\n".join(output).rstrip() + "\n"


# ---------------------------------------------------------------------------
# Écriture et programme principal
# ---------------------------------------------------------------------------

def write_markdown(command: str, markdown: str) -> Path:
    """Écrit le fichier Markdown dans le dossier de documentation."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    destination = OUTPUT_DIR / filename_for(command)
    destination.write_text(markdown, encoding="utf-8", newline="\n")
    return destination


def print_usage() -> None:
    program = SCRIPT_PATH.name
    print(
        f"Utilisation : python tools/{program} <commande>\n"
        f"Exemples   : python tools/{program} A\n"
        f"             python tools/{program} JM\n"
        f"             python tools/{program} ZD\n"
        f"Alias      : FQ, BANG, QUOTE, COUNTER, EQUALS, COLON, AT, BACKSLASH",
        file=sys.stderr,
    )


def main() -> int:
    if len(sys.argv) != 2:
        print_usage()
        return 2

    try:
        requested = normalize_command(sys.argv[1])
        source = locate_source()
        paragraphs = load_paragraphs(source)
        sections = collect_command_sections(paragraphs)

        if requested not in sections:
            available = ", ".join(sections.keys())
            print(
                f"Erreur : commande {requested!r} absente de DNB14A.docx.\n"
                f"Commandes détectées : {available}",
                file=sys.stderr,
            )
            return 1

        markdown = build_markdown(requested, sections[requested])
        destination = write_markdown(requested, markdown)

        print(f"Source  : {source}")
        print(f"Commande: {requested}")
        print(f"Sortie  : {destination}")
        return 0

    except (FileNotFoundError, ValueError) as error:
        print(f"Erreur : {error}", file=sys.stderr)
        return 1
    except Exception as error:  # Dernier filet de sécurité pour un outil CLI.
        print(f"Erreur inattendue : {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
